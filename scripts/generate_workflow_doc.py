from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SCIM-Inventory-Workflow-Guide.docx"

BLUE = "2563EB"
DARK_BLUE = "0B1F3A"
MID_BLUE = "123A72"
LIGHT_BLUE = "EAF2FF"
PALE_BLUE = "F4F7FB"
GREEN = "0F9F6E"
AMBER = "D97706"
RED = "DC2626"
SLATE = "64748B"
LINE = "D7E3F3"
WHITE = "FFFFFF"
INK = "10233F"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(SLATE)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 32, DARK_BLUE),
        ("Subtitle", 14, SLATE),
        ("Heading 1", 22, DARK_BLUE),
        ("Heading 2", 15, BLUE),
        ("Heading 3", 11.5, MID_BLUE),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)

    header = section.header.paragraphs[0]
    header.text = "SCIM-INVENTORY  /  WORKFLOW & PROCESS GUIDE"
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = rgb(BLUE)
    header.runs[0].font.letter_spacing = Pt(0.7)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = brand.add_run("SCIM-INVENTORY")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(14)
    run.font.color.rgb = rgb(BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Workflow & Process Guide")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Inventory, Procurement, Warehouse, Vendor, Logistics, and Recruitment Operations")

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_border(cell, "B8D1FA", "10")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(
        "A practical operating model for administrators, procurement teams, warehouse teams, vendors, recruiters, and auditors."
    )
    r.font.size = Pt(12)
    r.font.color.rgb = rgb(MID_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(80)
    r = meta.add_run(f"Version 1.0  |  {date.today().strftime('%B %d, %Y')}\nPrepared from the implemented application architecture")
    r.font.size = Pt(10)
    r.font.color.rgb = rgb(SLATE)
    doc.add_page_break()


def add_callout(doc: Document, title: str, text: str, tone: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE if tone == BLUE else "FFF7E8")
    set_cell_border(cell, tone, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"{title}\n")
    r.bold = True
    r.font.color.rgb = rgb(tone)
    r = p.add_run(text)
    r.font.color.rgb = rgb(INK)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        set_cell_shading(cell, DARK_BLUE)
        set_cell_border(cell, DARK_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = rgb(WHITE)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(value))
            r.font.size = Pt(8.7)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def load_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/seguisb.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE) -> None:
    draw.line([start, end], fill=f"#{color}", width=5)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - direction * 15, ey - 9), (ex - direction * 15, ey + 9)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 9, ey - direction * 15), (ex + 9, ey - direction * 15)]
    draw.polygon(points, fill=f"#{color}")


def create_flowchart(path: Path, title: str, rows: list[list[tuple[str, str]]], captions: list[str] | None = None) -> None:
    width = 1800
    margin_x = 80
    top = 165
    row_height = 255
    height = top + len(rows) * row_height + 105
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(38, True)
    subtitle_font = load_font(18, False)
    node_font = load_font(24, True)
    step_font = load_font(17, True)
    lane_font = load_font(20, True)
    legend_font = load_font(17, False)
    draw.rounded_rectangle((24, 24, width - 24, height - 24), radius=26, fill="#F7FAFF", outline="#D7E3F3", width=3)
    draw.text((margin_x, 55), title, font=title_font, fill="#0B1F3A")
    draw.text((margin_x, 108), "Follow the numbered path; diamonds indicate review or decision checkpoints.", font=subtitle_font, fill="#64748B")

    step_number = 1
    for row_index, nodes in enumerate(rows):
        y = top + row_index * row_height
        lane_width = 145 if captions else 0
        if row_index:
            draw.line((margin_x, y - 28, width - margin_x, y - 28), fill="#E1EAF6", width=2)
        if captions:
            draw.rounded_rectangle((margin_x, y + 33, margin_x + 125, y + 183), radius=18, fill="#0B1F3A")
            lane_lines = wrap_text(draw, captions[row_index], lane_font, 105)
            total = len(lane_lines) * 27
            for line_index, line in enumerate(lane_lines):
                bbox = draw.textbbox((0, 0), line, font=lane_font)
                draw.text((margin_x + 62 - (bbox[2] - bbox[0]) / 2, y + 108 - total / 2 + line_index * 27), line, font=lane_font, fill="white")
        available = width - margin_x * 2 - lane_width
        gap = 34
        node_width = int((available - gap * (len(nodes) - 1)) / len(nodes))
        x0 = margin_x + lane_width
        direction = 1 if row_index % 2 == 0 else -1
        node_bounds: list[tuple[int, int, int, int]] = []
        for node_index, (label, tone) in enumerate(nodes):
            visual_index = node_index if direction == 1 else len(nodes) - 1 - node_index
            x = x0 + visual_index * (node_width + gap)
            node_top = y + 33
            node_bottom = y + 183
            fill = {"blue": "#EAF2FF", "dark": "#123A72", "green": "#E7F8F2", "amber": "#FFF5E6", "red": "#FDECEC"}.get(tone, "#FFFFFF")
            outline = {"blue": "#2563EB", "dark": "#123A72", "green": "#0F9F6E", "amber": "#D97706", "red": "#DC2626"}.get(tone, "#94A3B8")
            text_color = "#FFFFFF" if tone == "dark" else "#10233F"
            if tone == "amber":
                points = [(x + node_width / 2, node_top), (x + node_width, (node_top + node_bottom) / 2), (x + node_width / 2, node_bottom), (x, (node_top + node_bottom) / 2)]
                shadow_points = [(px + 5, py + 7) for px, py in points]
                draw.polygon(shadow_points, fill="#D9E4F2")
                draw.polygon(points, fill=fill, outline=outline)
                draw.line(points + [points[0]], fill=outline, width=4, joint="curve")
                text_width = int(node_width * 0.62)
            else:
                draw.rounded_rectangle((x + 5, node_top + 7, x + node_width + 5, node_bottom + 7), radius=22, fill="#D9E4F2")
                draw.rounded_rectangle((x, node_top, x + node_width, node_bottom), radius=22, fill=fill, outline=outline, width=4)
                text_width = node_width - 42
            badge_fill = "#FFFFFF" if tone == "dark" else outline
            badge_text = outline if tone == "dark" else "#FFFFFF"
            badge_left = x + (int(node_width * 0.18) if tone == "amber" else 14)
            badge_center_x = badge_left + 18
            draw.ellipse((badge_left, node_top + 13, badge_left + 36, node_top + 49), fill=badge_fill)
            number_text = str(step_number)
            number_box = draw.textbbox((0, 0), number_text, font=step_font)
            draw.text((badge_center_x - (number_box[2] - number_box[0]) / 2, node_top + 31 - (number_box[3] - number_box[1]) / 2 - 2), number_text, font=step_font, fill=badge_text)
            lines = wrap_text(draw, label, node_font, text_width)
            total = len(lines) * 32
            for line_index, line in enumerate(lines):
                bbox = draw.textbbox((0, 0), line, font=node_font)
                draw.text((x + node_width / 2 - (bbox[2] - bbox[0]) / 2, (node_top + node_bottom) / 2 - total / 2 + line_index * 32 + 8), line, font=node_font, fill=text_color)
            node_bounds.append((x, node_top, x + node_width, node_bottom))
            step_number += 1
        for node_index in range(len(node_bounds) - 1):
            current = node_bounds[node_index]
            following = node_bounds[node_index + 1]
            center_y = int((current[1] + current[3]) / 2)
            if direction == 1:
                arrow(draw, (current[2] + 5, center_y), (following[0] - 5, center_y))
            else:
                arrow(draw, (current[0] - 5, center_y), (following[2] + 5, center_y))
        if row_index < len(rows) - 1:
            current_last = node_bounds[-1]
            start_x = int((current_last[0] + current_last[2]) / 2)
            next_count = len(rows[row_index + 1])
            next_available = width - margin_x * 2 - lane_width
            next_width = int((next_available - gap * (next_count - 1)) / next_count)
            next_direction = 1 if (row_index + 1) % 2 == 0 else -1
            next_visual_index = 0 if next_direction == 1 else next_count - 1
            next_x = margin_x + lane_width + next_visual_index * (next_width + gap)
            end_x = int(next_x + next_width / 2)
            mid_y = y + 218
            draw.line([(start_x, current_last[3] + 5), (start_x, mid_y), (end_x, mid_y)], fill=f"#{MID_BLUE}", width=5)
            arrow(draw, (end_x, mid_y), (end_x, y + row_height + 25), MID_BLUE)

    legend_y = height - 72
    legend_items = [("Activity", BLUE), ("Review / decision", AMBER), ("Confirmed outcome", GREEN), ("Controlled checkpoint", MID_BLUE)]
    legend_x = margin_x
    for label, color in legend_items:
        draw.rounded_rectangle((legend_x, legend_y, legend_x + 22, legend_y + 22), radius=5, fill=f"#{color}")
        draw.text((legend_x + 31, legend_y - 1), label, font=legend_font, fill="#475569")
        label_width = draw.textbbox((0, 0), label, font=legend_font)[2]
        legend_x += label_width + 82
    image.save(path, quality=95)


def add_diagram(doc: Document, image_path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.9))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = rgb(SLATE)


def add_contents(doc: Document) -> None:
    doc.add_heading("How to use this guide", level=1)
    doc.add_paragraph(
        "This document describes the operating workflows implemented in SCIM-Inventory. It is intended for process walkthroughs, user onboarding, project defense, testing, and future implementation planning."
    )
    add_callout(
        doc,
        "Core control principle",
        "Inventory quantities are changed through controlled stock movements—not by directly editing a balance. Every tenant-scoped operation is protected by membership, role checks, row-level security, and audit history.",
    )
    doc.add_heading("Contents", level=2)
    add_table(
        doc,
        ["Section", "Purpose"],
        [
            ["1. Operating model", "Actors, modules, ownership, and system boundaries"],
            ["2. Inventory and warehouse", "Products, locations, balances, movements, transfers, and tasks"],
            ["3. Procurement", "Requisition through purchase order and receiving"],
            ["4. Vendor lifecycle", "Onboarding, portal access, collaboration, shipment, invoice, and performance"],
            ["5. Recruitment and AI", "Public application, controlled stages, explainable ranking, and human review"],
            ["6. Notifications and communication", "In-app events, email notifications, and transaction messages"],
            ["7. Security and governance", "RBAC, MFA, tenant isolation, audit, documents, and reset controls"],
            ["8. Status and responsibility reference", "Operational status meanings and role ownership"],
        ],
        [2.0, 4.7],
    )


def build_document() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    with tempfile.TemporaryDirectory() as temp_dir_name:
        temp_dir = Path(temp_dir_name)

        master = temp_dir / "master.png"
        create_flowchart(
            master,
            "End-to-End Operating Model",
            [
                [("Demand or business need", "blue"), ("Requisition and sourcing", "blue"), ("Purchase order", "dark"), ("Vendor confirmation", "green")],
                [("Shipment and tracking", "blue"), ("Receipt and inspection", "amber"), ("Location stock update", "green"), ("Invoice and payment", "dark")],
            ],
        )
        add_diagram(doc, master, "Master workflow from demand signal to inventory and financial closure.")

        doc.add_heading("1. Operating model", level=1)
        doc.add_paragraph(
            "SCIM-Inventory is a multi-tenant operations platform. Each organization owns its products, warehouses, procurement records, vendors, recruitment records, documents, notifications, and audit trail. Users gain access through organization membership; vendor personnel use a separately scoped portal."
        )
        add_table(
            doc,
            ["Actor", "Primary responsibility", "Typical access"],
            [
                ["Owner", "Organization governance and final control", "All modules, membership, reset, audit"],
                ["Administrator", "Cross-functional operational administration", "Operations, vendors, recruitment; no ownership transfer"],
                ["Procurement manager / Buyer", "Sourcing, supplier, and purchase records", "Products, vendors, requisitions, RFQs, quotations, POs"],
                ["Warehouse manager / Operator", "Physical stock execution", "Locations, movements, receiving, putaway, picking, counts"],
                ["HR manager / Recruiter", "Hiring workflow and screening review", "Jobs, applicants, pipeline, AI screening"],
                ["Vendor user", "Supplier-side transaction collaboration", "Vendor-scoped orders, shipments, invoices, documents, messages"],
                ["Viewer", "Read-only operational visibility", "Permitted supply-chain records without mutation rights"],
            ],
            [1.45, 2.8, 2.45],
        )
        doc.add_heading("System navigation", level=2)
        add_bullets(
            doc,
            [
                "Overview provides real-time KPIs, service level, attention items, inbound orders, and recent movements.",
                "Inventory owns product master data, location-aware stock, immutable movements, transfers, and cycle counts.",
                "Warehouse owns receiving, putaway, picking, and physical location definitions.",
                "Procurement owns requisitions, RFQs, quotations, and purchase orders.",
                "Vendors provides onboarding, portal users, transactions, performance, documents, and conversation history.",
                "Recruitment manages job openings, public applications, the hiring pipeline, and assisted ranking.",
                "Administration manages users, roles, warehouses, assignments, audit history, and controlled workspace reset.",
            ],
        )

        doc.add_page_break()
        doc.add_heading("2. Inventory and warehouse workflow", level=1)
        inventory = temp_dir / "inventory.png"
        create_flowchart(
            inventory,
            "Location-Aware Inventory Flow",
            [
                [("Create product", "blue"), ("Create warehouse", "blue"), ("Create storage location", "blue"), ("Post receipt movement", "dark")],
                [("Validate role and scope", "amber"), ("Lock location balance", "amber"), ("Append immutable movement", "green"), ("Update on-hand balance", "green")],
                [("Receiving task", "blue"), ("Putaway confirmation", "blue"), ("Picking or issue", "dark"), ("Live dashboard refresh", "green")],
            ],
            ["SETUP", "POSTING", "EXECUTION"],
        )
        add_diagram(doc, inventory, "A quantity is always posted through a controlled movement and tied to a warehouse location.")
        doc.add_heading("Inventory rules", level=2)
        add_numbered(
            doc,
            [
                "Create the product master with its automatically generated SKU, category, unit of measure, and reorder point.",
                "Create a warehouse and one or more active physical locations such as receiving, storage, picking, staging, or quarantine.",
                "Post a movement by selecting the warehouse, location, product, movement type, quantity, and notes.",
                "The server validates permissions and confirms that the product, warehouse, and location are active and belong to the organization.",
                "A receipt or transfer-in increases stock; an issue or transfer-out decreases stock; an adjustment applies the signed quantity.",
                "The database prevents a transaction that would make location stock negative.",
                "The immutable movement is appended and the balance projection is updated in the same transaction.",
                "Realtime refresh updates the Overview, Stock, Movement, and Location screens.",
            ],
        )
        add_callout(
            doc,
            "Existing stock migration rule",
            "When location inventory is enabled, stock is automatically assigned only if a warehouse has exactly one active location. If multiple locations exist, historical stock remains labeled Unallocated until staff post a controlled location movement.",
            AMBER,
        )
        add_table(
            doc,
            ["Movement", "Balance effect", "Common use"],
            [
                ["Receipt", "+ absolute quantity", "Purchased or returned goods received"],
                ["Issue", "- absolute quantity", "Consumption, outbound fulfillment, or disposal"],
                ["Adjustment", "+/- entered quantity", "Approved correction following verification"],
                ["Transfer out", "- absolute quantity", "Dispatch from source location"],
                ["Transfer in", "+ absolute quantity", "Receipt into destination location"],
            ],
            [1.5, 1.8, 3.4],
        )
        doc.add_heading("Warehouse execution", level=2)
        add_table(
            doc,
            ["Workflow", "Typical statuses", "Completion evidence"],
            [
                ["Receiving", "Open → Assigned → In progress → Completed", "Received quantity, inspection, reference, date"],
                ["Putaway", "Open → Assigned → In progress → Completed", "Destination location and confirmed quantity"],
                ["Picking", "Open → Assigned → In progress → Completed", "Picked quantity and outbound context"],
                ["Cycle count", "Draft → Submitted → In progress → Completed", "Expected count, result, and approved variance"],
                ["Transfer", "Draft → Submitted → In progress → Completed", "Source, destination, quantity, and paired movements"],
            ],
            [1.25, 2.55, 2.9],
        )

        doc.add_page_break()
        doc.add_heading("3. Procurement workflow", level=1)
        procurement = temp_dir / "procurement.png"
        create_flowchart(
            procurement,
            "Procure-to-Receive Workflow",
            [
                [("Need identified", "blue"), ("Purchase requisition", "blue"), ("Review and approval", "amber"), ("Request for quotation", "dark")],
                [("Vendor quotations", "blue"), ("Commercial comparison", "amber"), ("Purchase order", "dark"), ("Vendor response", "green")],
                [("Processing and packing", "blue"), ("Shipment in transit", "blue"), ("Receipt and inspection", "amber"), ("PO completion", "green")],
            ],
            ["REQUEST", "SOURCE", "FULFILL"],
        )
        add_diagram(doc, procurement, "Procurement records retain references, organization scope, status, value, supplier, warehouse, due date, and audit history.")
        add_numbered(
            doc,
            [
                "A user creates a requisition manually or in response to low stock and operational demand.",
                "The requisition is submitted and reviewed according to role and organization controls.",
                "An RFQ is issued to qualified vendors, and quotations are captured for comparison.",
                "The buyer selects the commercial option and creates a purchase order with supplier, warehouse, amount, due date, and terms.",
                "The vendor confirms, proposes a revision, or rejects the purchase order through the portal.",
                "Accepted work progresses through processing, packing, ready for shipment, and transit.",
                "Warehouse staff receive and inspect delivered goods. Accepted quantity updates inventory; partial or rejected quantity creates an exception path.",
                "The purchase order is closed only after operational and document requirements are complete.",
            ],
        )
        doc.add_heading("Procurement controls", level=2)
        add_bullets(
            doc,
            [
                "References are generated automatically and displayed as read-only values.",
                "Organization and role checks are repeated on the server for every mutation.",
                "Purchase order acknowledgement preserves confirmed, revised, or rejected responses.",
                "Amounts use PHP as the current default currency.",
                "Changes create audit or workflow events; stock is never increased merely because a PO exists.",
            ],
        )

        doc.add_page_break()
        doc.add_heading("4. Vendor lifecycle and portal workflow", level=1)
        vendor = temp_dir / "vendor.png"
        create_flowchart(
            vendor,
            "Vendor Lifecycle",
            [
                [("Public application", "blue"), ("Credential review", "amber"), ("Approve vendor", "green"), ("Invite portal users", "dark")],
                [("Catalog and contacts", "blue"), ("PO acknowledgement", "blue"), ("Shipment updates", "blue"), ("Invoice submission", "dark")],
                [("Three-way match", "amber"), ("Payment record", "green"), ("Performance scorecard", "green"), ("Review or improvement plan", "dark")],
            ],
            ["ONBOARD", "TRANSACT", "EVALUATE"],
        )
        add_diagram(doc, vendor, "Vendor users are scoped to one organization and supplier. Internal users retain approval and governance responsibilities.")
        doc.add_heading("Phase A — onboarding and setup", level=2)
        add_bullets(
            doc,
            [
                "Vendor submits company, contact, service, location, delivery, certification, tax, and legal information.",
                "Internal staff review credentials and move onboarding through draft, submitted, under review, changes requested, approved, rejected, suspended, or archived.",
                "Approved vendors receive a unique vendor code and may be assigned categories, commercial terms, risk rating, catalog items, shipping rules, contacts, and addresses.",
                "An owner or authorized procurement user invites vendor personnel to the portal.",
            ],
        )
        doc.add_heading("Portal roles", level=2)
        add_table(
            doc,
            ["Vendor role", "Purpose", "Typical capabilities"],
            [
                ["Admin", "Vendor account administrator", "Manage vendor profile and portal users; broad vendor access"],
                ["Member", "General transaction user", "Catalog and routine purchase-order collaboration"],
                ["Finance", "Billing contact", "Invoices, financial documents, and payment visibility"],
                ["Logistics", "Shipment coordinator", "Shipment creation, milestones, tracking, and delivery updates"],
            ],
            [1.1, 2.0, 3.6],
        )
        doc.add_heading("Phase B — transaction collaboration", level=2)
        add_numbered(
            doc,
            [
                "Admin sends the purchase order digitally and the vendor receives it in the portal.",
                "Vendor responds Confirmed, Revised, or Rejected. Proposed price, quantity, date, and message remain traceable.",
                "Vendor creates a shipment linked to the purchase order, then posts milestones and tracking updates.",
                "Delivery creates internal and vendor notifications. Warehouse staff perform receipt and inspection.",
                "Vendor submits an invoice. The system compares invoice, PO, and received quantity and records matched or exception status.",
                "Authorized staff approve and record payment according to agreed terms.",
            ],
        )
        doc.add_heading("Performance model", level=2)
        add_table(
            doc,
            ["Indicator", "Meaning", "Weight in overall score"],
            [
                ["On-time delivery", "Delivered by expected arrival or due date", "30%"],
                ["Fulfillment", "Received quantity compared with ordered quantity", "30%"],
                ["Quality", "Accepted quantity compared with inspected quantity", "20%"],
                ["Accuracy", "Matched invoices compared with total matches", "15%"],
                ["Response", "Acknowledgement speed, normalized across 48 hours", "5%"],
            ],
            [1.55, 3.55, 1.6],
        )

        doc.add_page_break()
        doc.add_heading("5. Recruitment and AI-assisted screening", level=1)
        recruitment = temp_dir / "recruitment.png"
        create_flowchart(
            recruitment,
            "Controlled Recruitment Pipeline",
            [
                [("Applied", "blue"), ("Initial review", "amber"), ("AI screening", "dark"), ("Interview", "blue"), ("Hired", "green")],
            ],
        )
        add_diagram(doc, recruitment, "The simplified hiring pipeline. Rejected and withdrawn outcomes are stored separately without creating extra active stages.")
        add_numbered(
            doc,
            [
                "HR publishes an open job with description, required and preferred skills, minimum experience, and education requirement.",
                "A candidate applies through the public careers page, gives screening consent, and uploads a PDF or DOCX resume.",
                "The application receives an automatic reference and enters Applied.",
                "A recruiter checks completeness and moves the candidate to Initial review.",
                "Authorized staff run AI screening for the selected job. The service receives identity-minimized resume content and structured qualifications.",
                "The system records scores, rank, matched skills, missing skills, model name, model version, weights, and explanations.",
                "A human recruiter reviews the evidence and decides whether to move the candidate to Interview and then Hired.",
                "Stage changes can trigger applicant email notifications when mail delivery is configured.",
            ],
        )
        doc.add_heading("How the match score is calculated", level=2)
        add_table(
            doc,
            ["Signal", "Default weight", "Calculation"],
            [
                ["Semantic similarity", "40%", "TF-IDF unigrams/bigrams plus cosine similarity between job and resume text"],
                ["Skills", "35%", "80% required-skill coverage plus 20% preferred-skill coverage"],
                ["Experience", "15%", "Applicant years divided by required years, capped at 100%"],
                ["Education", "10%", "Applicant education level compared with requested level, capped at 100%"],
            ],
            [1.55, 1.2, 4.0],
        )
        add_callout(
            doc,
            "Responsible-use rule",
            "The ranking supports recruiter review; it does not make the hiring decision. Name, email, phone, age, gender, ethnicity, and address are not supplied to the model. A human must evaluate evidence, accommodations, and job relevance before changing a hiring stage.",
        )

        doc.add_page_break()
        doc.add_heading("6. Notifications, messaging, and delivery updates", level=1)
        comms = temp_dir / "communications.png"
        create_flowchart(
            comms,
            "Event-to-Notification Flow",
            [
                [("Business event", "blue"), ("Create workflow event", "blue"), ("Create scoped notification", "dark"), ("Realtime bell update", "green")],
                [("User opens notification", "blue"), ("Mark as read", "green"), ("Open linked record", "dark"), ("Continue transaction", "blue")],
            ],
            ["PUBLISH", "ACT"],
        )
        add_diagram(doc, comms, "Notifications are organization-scoped, role-protected, published in real time, and linked to the related dashboard record.")
        add_table(
            doc,
            ["Trigger", "Audience", "Result"],
            [
                ["New vendor application", "Internal", "Bell notification links to vendor review"],
                ["Vendor document submitted", "Internal", "Verification work becomes visible"],
                ["Shipment created or updated", "Internal and/or vendor", "Tracking and exception status is synchronized"],
                ["Invoice submitted", "Internal", "Finance is prompted to perform three-way matching"],
                ["Recruitment stage changed", "Applicant", "Configured email communicates the new stage"],
                ["Transaction message", "Internal or vendor conversation", "Conversation remains linked to PO, shipment, invoice, or general context"],
            ],
            [1.85, 1.4, 3.55],
        )
        add_bullets(
            doc,
            [
                "The bell displays the exact unread count and the 20 newest items.",
                "Users can mark one item or all visible-scope items as read.",
                "Opening a notification marks it read and routes to a validated internal dashboard URL.",
                "Realtime database events refresh the shell without requiring a manual page reload.",
                "Email delivery is separate from in-app notifications and depends on the configured SMTP provider.",
            ],
        )

        doc.add_page_break()
        doc.add_heading("7. Security and governance", level=1)
        security = temp_dir / "security.png"
        create_flowchart(
            security,
            "Authenticated Request Control",
            [
                [("Sign in", "blue"), ("Verify session", "blue"), ("MFA / AAL2 check", "amber"), ("Resolve membership", "dark")],
                [("Check organization role", "amber"), ("Apply RLS policy", "amber"), ("Execute transaction", "green"), ("Append audit event", "dark")],
            ],
            ["IDENTITY", "AUTHORIZE"],
        )
        add_diagram(doc, security, "UI visibility is a convenience; server authorization and database RLS are the enforcement boundaries.")
        add_table(
            doc,
            ["Control", "Implementation", "Risk reduced"],
            [
                ["Tenant isolation", "organization_id plus forced row-level security", "Cross-organization data access"],
                ["Role-based access", "Membership roles and permission checks", "Unauthorized business actions"],
                ["Vendor isolation", "Vendor-user membership scoped to supplier", "Access to another vendor’s transactions"],
                ["MFA", "TOTP enrollment, challenge, and AAL2 checks", "Account takeover for privileged roles"],
                ["Private files", "Private storage buckets and short-lived signed URLs", "Public exposure of documents and resumes"],
                ["Immutable records", "Movement ledger and controlled RPC posting", "Silent inventory-history alteration"],
                ["Audit history", "Actor, action, entity, metadata, and timestamp", "Untraceable governance changes"],
            ],
            [1.45, 3.0, 2.35],
        )
        doc.add_heading("Workspace data reset", level=2)
        add_numbered(
            doc,
            [
                "Only an organization owner can request a reset.",
                "The session must satisfy the privileged authentication level.",
                "The owner must type the exact organization name to confirm intent.",
                "An advisory lock prevents concurrent resets for the same organization.",
                "Operational, vendor, recruitment, inventory, document metadata, warehouse, and related scoped records are deleted transactionally.",
                "The organization, memberships, profiles, and pre-existing audit history remain.",
                "Trigger-generated deletion noise is collapsed into one organization.data_reset governance event containing the deleted-record count.",
            ],
        )
        add_callout(
            doc,
            "Destructive operation",
            "A reset produces an empty operational workspace and cannot be undone from the application. Export or back up required records before execution.",
            RED,
        )

        doc.add_page_break()
        doc.add_heading("8. Status and responsibility reference", level=1)
        add_table(
            doc,
            ["Domain", "Primary statuses", "Business owner"],
            [
                ["Products", "Active, Inactive", "Procurement / Warehouse management"],
                ["Inventory operation", "Draft, Submitted, In progress, Completed, Cancelled", "Warehouse management"],
                ["Warehouse task", "Open, Assigned, In progress, Completed, Cancelled", "Warehouse management / Operator"],
                ["Procurement", "Draft, Submitted, Approved, Sent, Received, Completed, Rejected, Cancelled", "Procurement manager / Buyer"],
                ["Vendor onboarding", "Draft, Submitted, Under review, Changes requested, Approved, Rejected, Suspended, Archived", "Owner / Procurement"],
                ["Shipment", "Planned, Processing, Packing, Ready, Dispatched, In transit, Arrived, Delivered, Exception, Cancelled", "Vendor logistics / Internal logistics"],
                ["Goods receipt", "Draft, Inspecting, Partial, Accepted, Rejected, Completed", "Warehouse management"],
                ["Invoice", "Draft, Submitted, Matching, Exception, Approved, Scheduled, Paid, Overdue, Rejected, Cancelled", "Finance-authorized vendor user / Internal procurement"],
                ["Recruitment", "Applied, Initial review, AI screening, Interview, Hired", "HR manager / Recruiter"],
            ],
            [1.35, 3.95, 1.5],
        )
        doc.add_heading("Recommended daily operating rhythm", level=2)
        add_table(
            doc,
            ["When", "Activity", "Expected result"],
            [
                ["Start of day", "Review Overview attention queue, service level, inbound work, and notifications", "Priority exceptions assigned"],
                ["During operations", "Post receipts/issues, complete tasks, update shipments, and resolve vendor messages", "Current stock and transaction status"],
                ["Before close", "Check negative/zero availability risks, overdue POs, invoice exceptions, and incomplete receipts", "Controlled carry-over list"],
                ["Weekly", "Review supplier history, open improvement plans, aging invoices, and recruitment pipeline", "Management follow-up"],
                ["Monthly / quarterly", "Refresh vendor scorecards and audit role, membership, warehouse, and security controls", "Governance evidence and performance action"],
            ],
            [1.25, 3.75, 1.8],
        )

        doc.add_heading("Process acceptance checklist", level=2)
        add_bullets(
            doc,
            [
                "Every record is scoped to the correct organization and, where applicable, warehouse, location, supplier, or job.",
                "The acting role is authorized for the requested transition.",
                "References are generated automatically and remain visible throughout the workflow.",
                "Inventory quantity changes are supported by immutable movement records.",
                "External or financial changes have the necessary document and status evidence.",
                "Exceptions are visible in the attention queue, notification center, workflow history, or audit log.",
                "A human remains responsible for approvals, supplier action, inventory corrections, and hiring decisions.",
            ],
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SCIM-Inventory Workflow & Process Guide"
    doc.core_properties.subject = "End-to-end operational workflows and flowcharts"
    doc.core_properties.author = "SCIM-Inventory"
    doc.core_properties.keywords = "inventory, procurement, warehouse, vendor, recruitment, workflow"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
